"""Báo cáo Excel "Lịch nạp LNG" — mẫu người dùng chốt 06/10/2026.

Từ ảnh mẫu (28/09 → 02/01): bỏ cột "Công ty TNHH Nhiên Liệu Xanh" và "Cảng Cái
Mép", lượng đổi từ tấn sang m³, giữ khách hàng + ngày; ngày nghỉ 24/11 nằm giữa hai
lần giao 21/11 và 27/11 thì ghi chú ở dòng giao TRƯỚC nó (21/11). File ba sheet:
Lịch nạp, Nhật ký nạp thật, Tiêu thụ theo ngày.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.api.routers import plan as plan_router
from app.api.schemas import PlanExportIn, PlanSettingsIn
from app.config import Settings
from app.db.models import Telemetry, Terminal
from app.repositories import plan_readings as pr_repo
from app.services import plan_export as px

VN = ZoneInfo("Asia/Ho_Chi_Minh")
PSN = "YKH-TANK-01"
TODAY = date(2026, 10, 6)


def _wb(content: bytes | memoryview):
    return load_workbook(io.BytesIO(bytes(content)))


def _actual(d: date, hh: int, before: float, after: float) -> px.ActualRefill:
    return px.ActualRefill(at=datetime(d.year, d.month, d.day, hh, 29, tzinfo=VN),
                           before_m3=before, after_m3=after)


# ------------------------------------------------------------------- ghép lịch
def test_ngay_nghi_ghi_chu_o_dong_giao_TRUOC_no() -> None:
    """Đúng ví dụ của người dùng: nghỉ 24/11, giao 21/11 và 27/11."""
    rows = px.build_schedule(
        [px.PlannedRow(date(2026, 11, 21), 39.9), px.PlannedRow(date(2026, 11, 27), 39.9)],
        [], [date(2026, 11, 24)],
        from_day=date(2026, 11, 1), to_day=date(2026, 11, 30), today=TODAY,
    )
    assert [(r.day, r.notes) for r in rows] == [
        (date(2026, 11, 21), ["Nghỉ 24/11 (Thứ Ba)"]),
        (date(2026, 11, 27), []),
    ]


def test_qua_khu_la_so_do_tuong_lai_la_ke_hoach() -> None:
    rows = px.build_schedule(
        [
            px.PlannedRow(date(2026, 9, 29), 40.0),               # trước hôm nay: bỏ
            px.PlannedRow(date(2026, 10, 2), 37.35, forced=True),  # đã có xe thật: bỏ
            px.PlannedRow(date(2026, 10, 8), 39.9, forced=True),
        ],
        [_actual(date(2026, 9, 28), 14, 6.05, 41.36), _actual(date(2026, 10, 2), 9, 12.0, 50.0)],
        [],
        from_day=date(2026, 9, 27), to_day=date(2026, 10, 31), today=TODAY,
    )
    assert [(r.day, round(r.m3, 2), r.actual) for r in rows] == [
        (date(2026, 9, 28), 35.31, True),
        (date(2026, 10, 2), 38.0, True),
        (date(2026, 10, 8), 39.9, False),
    ]
    assert rows[0].notes == ["Đã nạp thực tế lúc 14:29"]
    assert rows[2].notes == ["Nạp chỉ định"]


def test_ngay_nghi_truoc_moi_lan_giao_khong_bi_mat() -> None:
    rows = px.build_schedule(
        [px.PlannedRow(date(2026, 10, 9), 40.0)], [], [date(2026, 10, 7)],
        from_day=date(2026, 10, 6), to_day=date(2026, 10, 31), today=TODAY,
    )
    assert rows[0].notes == ["Nghỉ 07/10 (Thứ Tư)"]


# ------------------------------------------------------------------- workbook
def _meta(**kw: object) -> px.ReportMeta:
    base: dict[str, object] = {
        "tank_name": "Bồn LNG - YKH-TANK-01", "psn": PSN, "customer": "Yokohama",
        "from_day": date(2026, 9, 27), "to_day": date(2026, 12, 31),
        "generated_at": datetime(2026, 10, 6, 9, 0),
        "basis": "Kế hoạch tính theo: tiêu thụ 6 m³/ngày",
    }
    base.update(kw)
    return px.ReportMeta(**base)  # type: ignore[arg-type]


def _report() -> bytes:
    sched = px.build_schedule(
        [px.PlannedRow(date(2026, 10, 8), 39.9), px.PlannedRow(date(2026, 11, 21), 39.9)],
        [_actual(date(2026, 9, 28), 14, 6.05, 41.36)], [date(2026, 11, 24)],
        from_day=date(2026, 9, 27), to_day=date(2026, 12, 31), today=TODAY,
    )
    days = [
        px.DayRow(date(2026, 9, 27), 15.4, 8.89, 0.0),
        px.DayRow(date(2026, 9, 28), 8.89, 38.86, 35.31),
        px.DayRow(date(2026, 9, 29), None, None, 0.0, note="Không có số đo"),
    ]
    return px.build_report(sched, [_actual(date(2026, 9, 28), 14, 6.05, 41.36)], days, _meta())


def test_ba_sheet_dung_ten_dung_thu_tu() -> None:
    assert _wb(_report()).sheetnames == ["Lịch nạp", "Nhật ký nạp thật", "Tiêu thụ theo ngày"]


def test_sheet_lich_nap_theo_mau_m3_khong_co_cot_bi_gach() -> None:
    ws = _wb(_report())["Lịch nạp"]
    assert ws["A1"].value == "LỊCH NẠP LNG"
    assert "Yokohama" in ws["A2"].value and "YKH-TANK-01" in ws["A2"].value
    head = [c.value for c in ws[6]]
    assert head == ["STT", "Khách hàng", "Ngày nạp", "Thứ", "Lượng (m³)", "Trạng thái", "Ghi chú"]
    flat = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Nhiên Liệu Xanh" not in flat and "Cái Mép" not in flat and "tấn" not in flat
    r1 = [c.value for c in ws[7]]
    assert r1[1] == "Yokohama" and r1[2].date() == date(2026, 9, 28) and r1[3] == "Thứ Hai"
    assert r1[4] == pytest.approx(35.31) and r1[5] == "Đã nạp"
    assert ws.cell(7, 3).number_format == "dd/mm/yyyy"
    r3 = [c.value for c in ws[9]]
    assert r3[2].date() == date(2026, 11, 21) and r3[5] == "Kế hoạch"
    assert r3[6] == "Nghỉ 24/11 (Thứ Ba)"
    assert ws.cell(10, 1).value == "Tổng cộng" and ws.cell(10, 5).value == "=SUM(E7:E9)"
    assert ws.cell(11, 5).value == '=SUMIF(F7:F9,"Đã nạp",E7:E9)'
    assert ws.freeze_panes == "A7" and ws.print_title_rows == "$6:$6"


def test_sheet_tieu_thu_tinh_dung_va_co_bieu_do() -> None:
    ws = _wb(_report())["Tiêu thụ theo ngày"]
    head_row = next(r for r in range(1, 10) if ws.cell(r, 1).value == "Ngày")
    first = head_row + 1
    assert ws.cell(first + 1, 5).value == pytest.approx(35.31)
    # 8.89 + 35.31 − 38.86
    assert ws.cell(first + 1, 6).value == pytest.approx(5.34)
    assert ws.cell(first + 2, 3).value is None and ws.cell(first + 2, 7).value == "Không có số đo"
    assert len(ws._charts) == 2


def test_chu_nguoi_go_khong_bao_gio_thanh_cong_thuc() -> None:
    sched = px.build_schedule([px.PlannedRow(date(2026, 10, 8), 1.0)], [], [],
                              from_day=date(2026, 10, 1), to_day=date(2026, 10, 31), today=TODAY)
    ws = _wb(px.build_report(sched, [], [], _meta(customer="=1+1")))["Lịch nạp"]
    cell = next(c for row in ws.iter_rows() for c in row if c.value == "=1+1")
    assert cell.data_type == "s"


def test_ky_rong_van_ra_file_hop_le() -> None:
    wb = _wb(px.build_report([], [], [], _meta()))
    assert wb["Nhật ký nạp thật"].cell(7, 1).value == "Không có lần nạp nào trong kỳ."


# ------------------------------------------------------------------- schema
def test_dong_ngoai_ky_ky_nguoc_ky_qua_dai_bi_tu_choi() -> None:
    with pytest.raises(ValidationError):
        PlanExportIn(from_day=date(2026, 10, 1), to_day=date(2026, 10, 31),
                     rows=[{"day": date(2026, 11, 1), "m3": 40}])
    with pytest.raises(ValidationError):
        PlanExportIn(from_day=date(2026, 10, 31), to_day=date(2026, 10, 1), rows=[])
    with pytest.raises(ValidationError):
        PlanExportIn(from_day=date(2026, 1, 1), to_day=date(2027, 1, 2), rows=[])
    with pytest.raises(ValidationError):
        PlanExportIn(from_day=date(2026, 10, 1), to_day=date(2026, 10, 31),
                     rows=[{"day": date(2026, 10, 3), "m3": 0}])


def test_chuoi_rong_la_xoa() -> None:
    b = PlanSettingsIn(customer_name="   ")
    assert b.customer_name is None and "customer_name" in b.model_fields_set


# ------------------------------------------------------------------- route + DB
@pytest.mark.db
def test_route_lay_lan_nap_that_va_tieu_thu_tu_so_do(session: Session) -> None:
    now = datetime.now(VN)
    today = now.date()
    t = Terminal(psn=PSN, name="Bồn LNG - YKH-TANK-01", capacity_l=Decimal("60000"))
    session.add(t)
    session.flush()
    start = datetime(today.year, today.month, today.day, tzinfo=VN) - timedelta(days=3)
    vol, at = 30000.0, start
    while at < now - timedelta(minutes=30):
        if at == start + timedelta(days=1, hours=14):
            vol += 35000.0                                    # xe tới
        session.add(Telemetry(terminal_id=t.id, psn=PSN, sampled_at=at,
                              volume_l=Decimal(str(round(vol, 1))), source="tst", raw_payload={}))
        vol -= 120.0
        at += timedelta(minutes=30)
    session.flush()
    pr_repo.save_settings(session, PSN, {"customer_name": "Yokohama"}, by="t")
    rest_day = today + timedelta(days=6)
    pr_repo.set_flags(session, PSN, rest_day, rest=True, forced=False, no_delivery=True)

    body = PlanExportIn(
        from_day=start.date(), to_day=today + timedelta(days=20),
        basis="tiêu thụ 5.76 m³/ngày",
        rows=[{"day": today + timedelta(days=4), "m3": "40", "forced": False}],
    )
    resp = plan_router.export_refills_xlsx(PSN, body, session, Settings(), "t")  # type: ignore[arg-type]
    assert resp.headers["content-disposition"].startswith('attachment; filename="lich_nap_YKH-TANK-01_')
    wb = _wb(resp.body)

    ws = wb["Lịch nạp"]
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=7, max_row=8)]
    assert data[0][5] == "Đã nạp" and data[0][2].date() == (start + timedelta(days=1)).date()
    assert data[0][4] == pytest.approx(35.0, abs=0.5)
    assert data[1][5] == "Kế hoạch" and data[1][1] == "Yokohama"
    assert f"Nghỉ {rest_day:%d/%m}" in (data[1][6] or "")   # dòng giao trước ngày nghỉ

    log = wb["Nhật ký nạp thật"]
    assert log.cell(7, 6).value == pytest.approx(35.0, abs=0.5)

    days = wb["Tiêu thụ theo ngày"]
    first_use = days.cell(7, 6).value
    assert first_use == pytest.approx(48 * 0.12, abs=0.3)  # 48 lần × 120 L
