"""Kế hoạch nạp: số đo thật cho ngày đã qua, lịch không giao Chủ Nhật/ngày nghỉ,
khung khuyến nghị khớp bảng, đánh dấu ngày nghỉ được LƯU, xuất Excel."""

from __future__ import annotations

import io
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from tests.e2e.conftest import go
from tests.e2e.seed import CUSTOMER, YKH


def _open_tank(page: Page) -> None:
    go(page, "planner")
    page.locator("#p-tank").select_option(YKH)
    expect(page.locator("#plan-anchor")).to_contain_text("số đo thật")
    expect(page.locator("#plan-body tr").first).to_be_visible()


def _rows(page: Page) -> list[dict[str, Any]]:
    return page.evaluate(
        """() => [...document.querySelectorAll('#plan-body tr')].map(tr => {
             const td = tr.querySelectorAll('td');
             return { day: td[0].textContent.trim(), wd: td[1].textContent.trim(),
                      lvl: td[2].textContent.trim(), act: td[6].textContent.trim(),
                      when: td[7].textContent.trim(), order: td[8].textContent.trim(),
                      key: td[4].querySelector('input').dataset.day,
                      rest: td[4].querySelector('input').checked,
                      locked: td[4].querySelector('input').disabled };
           })"""
    )


def test_ngay_da_qua_la_so_do_that_va_bi_khoa(app_page: Page, world: dict[str, Any]) -> None:
    _open_tank(app_page)
    start = world["today"] - timedelta(days=5)
    app_page.locator("#p-date").fill(start.isoformat())
    expect(app_page.locator("#plan-body tr").first).to_contain_text(f"{start:%d/%m}")
    expect(app_page.locator("#plan-body tr").first).to_contain_text("đo thật")
    rows = _rows(app_page)
    past = [r for r in rows if date.fromisoformat(r["key"]) < world["today"]]
    assert len(past) == 5 and all(r["locked"] for r in past)
    assert not any("Nạp LNG" in r["act"] for r in past), "bịa ra lần nạp trong quá khứ"


def test_lich_khong_giao_chu_nhat_va_ngay_nghi(app_page: Page, world: dict[str, Any]) -> None:
    _open_tank(app_page)
    rows = _rows(app_page)
    fills = [r for r in rows if "Nạp LNG" in r["act"]]
    assert fills, "kế hoạch 60 ngày phải có lần nạp"
    assert not [r for r in fills if r["wd"] == "Chủ Nhật"]
    rest = world["rest_day"].isoformat()
    rest_row = next(r for r in rows if r["key"] == rest)
    assert rest_row["rest"] and "Nạp LNG" not in rest_row["act"]


def test_khung_khuyen_nghi_khop_bang(app_page: Page) -> None:
    _open_tank(app_page)
    first = next(r for r in _rows(app_page) if "Nạp LNG" in r["act"])
    box = app_page.locator("#plan-assist")
    expect(box).to_contain_text(f"giao lúc {first['when'][:5]}")
    expect(box).to_contain_text(f"{float(first['order']):.2f} m³")


def test_danh_dau_ngay_nghi_duoc_luu_va_bo_duoc(app_page: Page, server: str) -> None:
    _open_tank(app_page)
    fill = next(r for r in _rows(app_page) if "Nạp LNG" in r["act"])
    box = app_page.locator(f'#plan-body input[data-kind="rest"][data-day="{fill["key"]}"]')
    with app_page.expect_response(lambda r: "/api/plan/flags/" in r.url and r.request.method == "PUT"):
        box.check()
    moved = next(r for r in _rows(app_page) if r["key"] == fill["key"])
    assert "Nạp LNG" not in moved["act"], "lần nạp phải dời khỏi ngày nghỉ"

    app_page.reload()
    _open_tank(app_page)
    box = app_page.locator(f'#plan-body input[data-kind="rest"][data-day="{fill["key"]}"]')
    expect(box).to_be_checked()                         # đã lưu ở server, không chỉ trong tab
    with app_page.expect_response(lambda r: "/api/plan/flags/" in r.url and r.request.method == "PUT"):
        box.uncheck()
    back = next(r for r in _rows(app_page) if r["key"] == fill["key"])
    assert "Nạp LNG" in back["act"]


def test_xuat_excel_tu_trang_ke_hoach(app_page: Page) -> None:
    _open_tank(app_page)
    expect(app_page.locator("#p-x-customer")).to_have_value(CUSTOMER)
    with app_page.expect_download() as dl:
        app_page.locator("#plan-export-btn").click()
    wb = load_workbook(io.BytesIO(Path(dl.value.path()).read_bytes()))
    assert wb.sheetnames == ["Lịch nạp", "Nhật ký nạp thật", "Tiêu thụ theo ngày"]
    ws = wb["Lịch nạp"]
    assert CUSTOMER in ws["A4"].value
    assert [c.value for c in ws[7]][:3] == ["STT", "Khách hàng", "Ngày nạp"]
    assert ws.cell(8, 6).value in ("Đã nạp", "Kế hoạch")
