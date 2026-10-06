"""Báo cáo: báo cáo Excel 3 sheet, báo cáo in, lịch giao, báo động nhà máy."""

from __future__ import annotations

import io
from datetime import timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from tests.e2e.conftest import go
from tests.e2e.seed import CUSTOMER, YKH

VN = ZoneInfo("Asia/Ho_Chi_Minh")


def _open(page: Page) -> None:
    go(page, "reports")
    page.locator("#rep-tank").select_option(YKH)
    expect(page.locator("#rep-customer")).to_have_value(CUSTOMER)


def test_khong_con_csv_tho(app_page: Page) -> None:
    go(app_page, "reports")
    expect(app_page.locator("#view-reports")).not_to_contain_text("CSV")
    expect(app_page.locator("#rep-xlsx")).to_be_visible()


def test_bao_cao_excel_ba_sheet_dung_so_lieu_that(app_page: Page, world: dict[str, Any]) -> None:
    _open(app_page)
    today = world["today"]
    frm, to = today - timedelta(days=20), today + timedelta(days=40)
    app_page.locator("#rep-from").fill(frm.isoformat())
    app_page.locator("#rep-to").fill(to.isoformat())
    with app_page.expect_download() as dl:
        app_page.locator("#rep-xlsx").click()
    wb = load_workbook(io.BytesIO(Path(dl.value.path()).read_bytes()))

    sched = wb["Lịch nạp"]
    body = [[c.value for c in row] for row in sched.iter_rows(min_row=8) if isinstance(row[0].value, int)]
    real = [r for r in body if r[5] == "Đã nạp"]
    planned = [r for r in body if r[5] == "Kế hoạch"]
    seeded = {a.astimezone(VN).date() for a in world["refills"] if a.astimezone(VN).date() >= frm}
    assert {r[2].date() for r in real} == seeded, (body, seeded)
    assert planned and all(r[2].date() >= today for r in planned)
    assert all(r[1] == CUSTOMER for r in body)
    assert all(r[3] != "Chủ Nhật" for r in planned)
    rest = world["rest_day"]
    assert any(f"Nghỉ {rest:%d/%m}" in (r[6] or "") for r in body), "thiếu ghi chú ngày nghỉ"
    total = next(row for row in sched.iter_rows(min_row=8) if row[0].value == "Tổng cộng")
    assert total[4].value == round(sum(r[4] for r in body), 2)        # SỐ, không công thức

    log = wb["Nhật ký nạp thật"]
    assert len([r for r in log.iter_rows(min_row=8) if isinstance(r[0].value, int)]) == len(real)
    days = wb["Tiêu thụ theo ngày"]
    n_days = len([r for r in days.iter_rows(min_row=8) if hasattr(r[0].value, "date")])
    assert n_days == (today - frm).days + 1

    # Mượn bộ tính của trang Kế hoạch rồi trả lại: ô bồn ở Kế hoạch không đổi.
    expect(app_page.locator("#p-tank")).to_have_value("")


def test_bao_cao_in_mo_tab_moi(app_page: Page) -> None:
    go(app_page, "reports")
    link = app_page.locator("#rp-open")
    expect(link).to_have_attribute("target", "_blank")
    with app_page.context.expect_page() as pg:
        link.click()
    report = pg.value
    report.wait_for_load_state()
    assert report.title().startswith("BC-LNG-"), report.title()
    expect(report.locator("h1").first).to_contain_text("Báo cáo giám sát bồn LNG")


def test_lich_giao_va_bao_dong_nha_may(app_page: Page) -> None:
    go(app_page, "reports")
    app_page.locator("#btn-trips").click()
    expect(app_page.locator("#trip-list")).to_contain_text("Chuyến 1")
    view = app_page.locator("#view-reports")
    expect(view).to_contain_text("6 dòng thô → 2 việc")
    expect(view).to_contain_text("5.96–14.9 m3")
