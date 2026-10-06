"""Báo cáo Excel "Lịch nạp LNG": một file, ba sheet, in vừa A4.

Mẫu là bảng người vận hành gửi đối tác (ảnh 30/09/2026, sửa lại 06/10/2026): chỉ
các ngày xe tới, lượng tính bằng m³ (không phải tấn), cột khách hàng, ngày, ghi
chú — và ngày nghỉ rơi giữa hai lần giao thì ghi chú ở dòng giao TRƯỚC nó.

1. **Lịch nạp** — sheet chính: lần nạp THẬT đã qua (số đo) + lần nạp KẾ HOẠCH.
2. **Nhật ký nạp thật** — mỗi lần xe đã tới, kèm thể tích trước/sau.
3. **Tiêu thụ theo ngày** — đầu ngày, cuối ngày, nạp, tiêu thụ; kèm biểu đồ.

Ai tính gì. Lịch KẾ HOẠCH do trang tính (cùng bộ tính với trang Kế hoạch, cùng ô
tích và số đo tay người dùng đang thấy) và gửi lên; một bản tính thứ hai ở server
là cách chắc chắn nhất để file lệch với màn hình. Phần THẬT — lần nạp, thể tích
từng ngày, ngày nghỉ — server tự đọc từ dữ liệu, không tin trang gửi lên.

Module này thuần: nhận dữ liệu đã chuẩn bị, trả bytes. Không đọc DB.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

WEEKDAY_VI = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ Nhật"]

# Bảng màu: xanh thương hiệu cho tiêu đề, xanh lá nhạt = đã xảy ra, vàng = tổng
# (giữ đúng thói quen của mẫu gốc), xám = ngày nghỉ / không có số đo.
_BRAND = "1F4E79"
_HEAD_FILL = PatternFill("solid", fgColor=_BRAND)
_DONE_FILL = PatternFill("solid", fgColor="E2F0D9")
_TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
_MUTED_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_M3 = "#,##0.00"
_DATE = "dd/mm/yyyy"
_DT = "dd/mm/yyyy hh:mm"


# --------------------------------------------------------------------------- #
# Dữ liệu vào
# --------------------------------------------------------------------------- #


@dataclass(frozen=True)
class PlannedRow:
    """Một lần nạp KẾ HOẠCH do trang tính."""

    day: date
    m3: float
    forced: bool = False


@dataclass(frozen=True)
class ActualRefill:
    """Một lần xe đã tới, suy từ số đo. ``at`` là giờ địa phương."""

    at: datetime
    before_m3: float
    after_m3: float

    @property
    def m3(self) -> float:
        return self.after_m3 - self.before_m3


@dataclass(frozen=True)
class DayRow:
    day: date
    open_m3: float | None
    close_m3: float | None
    refill_m3: float
    rest: bool = False
    note: str = ""

    @property
    def use_m3(self) -> float | None:
        if self.open_m3 is None or self.close_m3 is None:
            return None
        return self.open_m3 + self.refill_m3 - self.close_m3


@dataclass(frozen=True)
class ScheduleRow:
    day: date
    m3: float
    actual: bool
    notes: list[str] = field(default_factory=list)


@dataclass(frozen=True)
class ReportMeta:
    tank_name: str
    psn: str
    customer: str | None
    from_day: date
    to_day: date
    generated_at: datetime
    basis: str | None = None


def vn_day(d: date) -> str:
    return f"{d:%d/%m} ({WEEKDAY_VI[d.weekday()]})"


# --------------------------------------------------------------------------- #
# Ghép lịch: thật + kế hoạch + ghi chú ngày nghỉ
# --------------------------------------------------------------------------- #


def build_schedule(
    planned: list[PlannedRow],
    actual: list[ActualRefill],
    rest_days: list[date],
    *,
    from_day: date,
    to_day: date,
    today: date,
) -> list[ScheduleRow]:
    """Các dòng của sheet Lịch nạp, theo ngày tăng dần.

    * Ngày đã qua chỉ có lần nạp THẬT. Một lần nạp kế hoạch rơi vào ngày đã có
      xe thật thì bỏ — xe đã tới, lịch mô phỏng cho ngày đó không còn nghĩa — và
      lần nạp kế hoạch trước hôm nay cũng bỏ: quá khứ là số đo, không phải dự kiến.
    * Ngày nghỉ (người dùng đánh dấu) ghi chú ở dòng giao TRƯỚC nó (người dùng
      chọn 06/10/2026): đọc dòng giao là biết chuyến sau phải qua một ngày nghỉ.
      Không có dòng nào trước nó thì ghi ở dòng đầu tiên, không được làm rơi mất.
    """
    rows: dict[date, ScheduleRow] = {}
    for a in sorted(actual, key=lambda x: x.at):
        d = a.at.date()
        if not from_day <= d <= to_day:
            continue
        note = f"Đã nạp thực tế lúc {a.at:%H:%M}"
        if d in rows:   # hai xe cùng ngày: gộp, giữ cả hai giờ
            prev = rows[d]
            rows[d] = ScheduleRow(d, prev.m3 + a.m3, True, [*prev.notes, note])
        else:
            rows[d] = ScheduleRow(d, a.m3, True, [note])
    for p in planned:
        if not from_day <= p.day <= to_day or p.day < today or p.day in rows or p.m3 <= 0:
            continue
        rows[p.day] = ScheduleRow(p.day, p.m3, False, ["Nạp chỉ định"] if p.forced else [])

    ordered = [rows[d] for d in sorted(rows)]
    for r in sorted(d for d in rest_days if from_day <= d <= to_day):
        before = [s for s in ordered if s.day <= r]
        target = before[-1] if before else (ordered[0] if ordered else None)
        if target is not None:
            target.notes.append(f"Nghỉ {vn_day(r)}")
    return ordered


# --------------------------------------------------------------------------- #
# Dựng workbook
# --------------------------------------------------------------------------- #


def _text(cell: Cell, value: str | None) -> None:
    """Ghi CHUỖI, không bao giờ là công thức.

    openpyxl coi mọi chuỗi mở đầu bằng "=" là công thức. Tên khách hàng hay ghi chú
    là chữ người gõ vào, và một ghi chú "=HYPERLINK(...)" không được phép chạy trên
    máy của người nhận file.
    """
    cell.value = value or ""
    cell.data_type = "s"


def _title_block(ws: Worksheet, title: str, meta: ReportMeta, ncols: int, extra: str | None = None) -> int:
    """Ba dòng đầu trang. Trả về số dòng của hàng tiêu đề bảng."""
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=15, color=_BRAND)
    ws.row_dimensions[1].height = 24

    ws.merge_cells(f"A2:{last}2")
    who = f"Bồn: {meta.tank_name} ({meta.psn})"
    if meta.customer:
        who += f"   ·   Khách hàng: {meta.customer}"
    _text(ws["A2"], who)
    ws["A2"].font = Font(bold=True, size=11)

    ws.merge_cells(f"A3:{last}3")
    line = (
        f"Kỳ: {meta.from_day:%d/%m/%Y} – {meta.to_day:%d/%m/%Y}   ·   "
        f"Xuất lúc {meta.generated_at:%d/%m/%Y %H:%M}"
    )
    _text(ws["A3"], line)
    ws["A3"].font = Font(size=10, color="595959")
    row = 4
    if extra:
        ws.merge_cells(f"A4:{last}4")
        _text(ws["A4"], extra)
        ws["A4"].font = Font(italic=True, size=9, color="595959")
        ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[4].height = 28
        row = 5
    return row + 1


def _header(ws: Worksheet, row: int, labels: list[str], widths: list[float]) -> None:
    for i, (lbl, w) in enumerate(zip(labels, widths, strict=True), start=1):
        c = ws.cell(row, i, lbl)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD_FILL
        c.border = _BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row + 1, 1)
    ws.print_title_rows = f"{row}:{row}"


def _print_setup(ws: Worksheet, landscape: bool = False) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "Trang &P / &N"
    ws.oddFooter.center.size = 8


def _row_style(ws: Worksheet, row: int, ncols: int, fill: PatternFill | None) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row, col)
        c.border = _BOX
        if fill is not None:
            c.fill = fill
        if not c.alignment.horizontal:
            c.alignment = Alignment(vertical="center")


def _sheet_schedule(ws: Worksheet, rows: list[ScheduleRow], meta: ReportMeta) -> None:
    ws.title = "Lịch nạp"
    labels = ["STT", "Khách hàng", "Ngày nạp", "Thứ", "Lượng (m³)", "Trạng thái", "Ghi chú"]
    n = len(labels)
    head = _title_block(ws, "LỊCH NẠP LNG", meta, n, meta.basis)
    _header(ws, head, labels, [6, 24, 13, 11, 13, 12, 46])

    r = head
    for i, s in enumerate(rows, start=1):
        r = head + i
        ws.cell(r, 1, i).alignment = Alignment(horizontal="center")
        _text(ws.cell(r, 2), meta.customer or meta.tank_name)
        c = ws.cell(r, 3, s.day)
        c.number_format = _DATE
        c.alignment = Alignment(horizontal="center")
        _text(ws.cell(r, 4), WEEKDAY_VI[s.day.weekday()])
        ws.cell(r, 5, round(s.m3, 2)).number_format = _M3
        _text(ws.cell(r, 6), "Đã nạp" if s.actual else "Kế hoạch")
        ws.cell(r, 6).alignment = Alignment(horizontal="center")
        _text(ws.cell(r, 7), "; ".join(s.notes))
        ws.cell(r, 7).alignment = Alignment(wrap_text=True, vertical="center")
        _row_style(ws, r, n, _DONE_FILL if s.actual else None)

    first, last = head + 1, max(head + 1, r)
    t = r + 1
    ws.merge_cells(start_row=t, start_column=1, end_row=t, end_column=4)
    ws.cell(t, 1, "Tổng cộng").alignment = Alignment(horizontal="center")
    # Công thức chứ không ghi số: người nhận hay sửa tay một dòng, và tổng phải theo.
    ws.cell(t, 5, f"=SUM(E{first}:E{last})" if rows else 0).number_format = _M3
    ws.cell(t, 6, f'=COUNTA(F{first}:F{last})&" lần"' if rows else "0 lần")
    ws.cell(t, 6).alignment = Alignment(horizontal="center")
    for col in range(1, n + 1):
        c = ws.cell(t, col)
        c.font = Font(bold=True)
        c.fill = _TOTAL_FILL
        c.border = _BOX

    # Tách tổng thành "đã nạp" và "kế hoạch" — câu hỏi đầu tiên của người đọc.
    for k, (lbl, key) in enumerate((("Trong đó đã nạp", "Đã nạp"), ("Trong đó kế hoạch", "Kế hoạch")), start=1):
        rr = t + k
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=4)
        ws.cell(rr, 1, lbl).alignment = Alignment(horizontal="right")
        ws.cell(rr, 1).font = Font(italic=True, color="595959")
        if rows:
            ws.cell(rr, 5, f'=SUMIF(F{first}:F{last},"{key}",E{first}:E{last})')
            ws.cell(rr, 6, f'=COUNTIF(F{first}:F{last},"{key}")&" lần"')
        else:
            ws.cell(rr, 5, 0)
            ws.cell(rr, 6, "0 lần")
        ws.cell(rr, 5).number_format = _M3
        ws.cell(rr, 6).alignment = Alignment(horizontal="center")
        if key == "Đã nạp":
            for col in (5, 6):
                ws.cell(rr, col).fill = _DONE_FILL

    _print_setup(ws)


def _sheet_refills(ws: Worksheet, refills: list[ActualRefill], meta: ReportMeta) -> None:
    ws.title = "Nhật ký nạp thật"
    labels = ["STT", "Thời điểm", "Thứ", "Trước khi nạp (m³)", "Sau khi nạp (m³)", "Lượng nạp (m³)"]
    n = len(labels)
    head = _title_block(
        ws, "NHẬT KÝ NẠP THẬT", meta, n,
        "Hệ thống tự nhận diện lần nạp từ bước tăng của thể tích đo được — không nhập tay.",
    )
    _header(ws, head, labels, [6, 18, 11, 16, 16, 15])
    r = head
    for i, a in enumerate(sorted(refills, key=lambda x: x.at), start=1):
        r = head + i
        ws.cell(r, 1, i).alignment = Alignment(horizontal="center")
        c = ws.cell(r, 2, a.at.replace(tzinfo=None))
        c.number_format = _DT
        c.alignment = Alignment(horizontal="center")
        _text(ws.cell(r, 3), WEEKDAY_VI[a.at.weekday()])
        ws.cell(r, 4, round(a.before_m3, 2)).number_format = _M3
        ws.cell(r, 5, round(a.after_m3, 2)).number_format = _M3
        ws.cell(r, 6, round(a.m3, 2)).number_format = _M3
        _row_style(ws, r, n, None)
    if not refills:
        ws.merge_cells(start_row=head + 1, start_column=1, end_row=head + 1, end_column=n)
        _text(ws.cell(head + 1, 1), "Không có lần nạp nào trong kỳ.")
        ws.cell(head + 1, 1).font = Font(italic=True, color="595959")
        r = head + 1
    t = r + 1
    ws.merge_cells(start_row=t, start_column=1, end_row=t, end_column=5)
    ws.cell(t, 1, f"Tổng cộng · {len(refills)} lần").alignment = Alignment(horizontal="center")
    ws.cell(t, 6, f"=SUM(F{head + 1}:F{r})" if refills else 0).number_format = _M3
    for col in range(1, n + 1):
        c = ws.cell(t, col)
        c.font = Font(bold=True)
        c.fill = _TOTAL_FILL
        c.border = _BOX
    _print_setup(ws)


def _chart_style(chart: LineChart | BarChart, title: str, cats: Reference) -> None:
    """Một kiểu cho mọi biểu đồ: MỘT màu, có số trên trục, không chú giải thừa.

    openpyxl 3.1 mặc định ẩn cả hai trục (``delete=True``) và tô mỗi điểm một màu
    (``varyColors``): biểu đồ ra không có một con số nào và trông như cầu vồng
    (đo bằng Excel xuất PDF, 06/10/2026).
    """
    chart.title = title
    chart.title.overlay = False     # tiêu đề nằm TRÊN vùng vẽ, không đè lưới
    chart.height, chart.width = 7.5, 17
    chart.set_categories(cats)
    chart.legend = None
    chart.varyColors = False
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.number_format = "dd/mm"
    chart.y_axis.number_format = "0"
    chart.y_axis.majorGridlines.spPr = None


def _sheet_days(ws: Worksheet, days: list[DayRow], meta: ReportMeta) -> None:
    ws.title = "Tiêu thụ theo ngày"
    labels = ["Ngày", "Thứ", "Đầu ngày (m³)", "Cuối ngày (m³)", "Nạp trong ngày (m³)",
              "Tiêu thụ (m³)", "Ghi chú"]
    n = len(labels)
    head = _title_block(
        ws, "TIÊU THỤ THEO NGÀY", meta, n,
        "Tiêu thụ = đầu ngày + nạp trong ngày − cuối ngày. Đầu ngày là lần đo đầu tiên "
        "của ngày; cuối ngày là đầu ngày hôm sau. Chỉ gồm những ngày đã qua.",
    )
    _header(ws, head, labels, [13, 11, 14, 14, 16, 14, 28])
    r = head
    for i, d in enumerate(days, start=1):
        r = head + i
        c = ws.cell(r, 1, d.day)
        c.number_format = _DATE
        c.alignment = Alignment(horizontal="center")
        _text(ws.cell(r, 2), WEEKDAY_VI[d.day.weekday()])
        for col, v in ((3, d.open_m3), (4, d.close_m3), (5, d.refill_m3 or None), (6, d.use_m3)):
            cell = ws.cell(r, col, None if v is None else round(v, 2))
            cell.number_format = _M3
        _text(ws.cell(r, 7), d.note)
        fill = _DONE_FILL if d.refill_m3 > 0 else (_MUTED_FILL if d.rest or d.open_m3 is None else None)
        _row_style(ws, r, n, fill)
    if days:
        t = r + 1
        ws.merge_cells(start_row=t, start_column=1, end_row=t, end_column=4)
        ws.cell(t, 1, "Tổng cộng").alignment = Alignment(horizontal="center")
        ws.cell(t, 5, f"=SUM(E{head + 1}:E{r})").number_format = _M3
        ws.cell(t, 6, f"=SUM(F{head + 1}:F{r})").number_format = _M3
        for col in range(1, n + 1):
            c = ws.cell(t, col)
            c.font = Font(bold=True)
            c.fill = _TOTAL_FILL
            c.border = _BOX
        avg = t + 1
        ws.merge_cells(start_row=avg, start_column=1, end_row=avg, end_column=5)
        # Bình quân những ngày CÓ tiêu thụ — cùng cách tính "khi nhà máy chạy" của
        # trang Kế hoạch. Gộp cả ngày nghỉ (tiêu thụ 0) thì số bị kéo thấp và lệch
        # với mức tiêu thụ người vận hành đang lập kế hoạch.
        ws.cell(avg, 1, "Bình quân ngày nhà máy chạy (bỏ ngày tiêu thụ 0)").alignment = Alignment(horizontal="right")
        ws.cell(avg, 1).font = Font(italic=True, color="595959")
        ws.cell(avg, 6, f'=IFERROR(AVERAGEIF(F{head + 1}:F{r},">0"),"")').number_format = _M3

        # Hai biểu đồ cạnh bảng: thể tích đầu ngày (thấy được nhịp nạp) và tiêu thụ
        # mỗi ngày (thấy được ngày nghỉ, ngày chạy mạnh). Nhìn một lần là hiểu kỳ.
        cats = Reference(ws, min_col=1, min_row=head + 1, max_row=r)
        line = LineChart()
        line.add_data(Reference(ws, min_col=3, min_row=head, max_row=r), titles_from_data=True)
        _chart_style(line, "Thể tích đầu ngày (m³)", cats)
        s = line.series[0]
        s.smooth = False          # đường gãy: nạp là một bước nhảy, không phải đường cong
        s.marker.symbol = "circle"
        s.marker.size = 5
        s.marker.graphicalProperties.solidFill = _BRAND
        s.marker.graphicalProperties.line.solidFill = _BRAND
        s.graphicalProperties.line.solidFill = _BRAND
        s.graphicalProperties.line.width = 22000
        ws.add_chart(line, f"{get_column_letter(n + 2)}{head}")
        bar = BarChart()
        bar.add_data(Reference(ws, min_col=6, min_row=head, max_row=r), titles_from_data=True)
        _chart_style(bar, "Tiêu thụ mỗi ngày (m³)", cats)
        bar.gapWidth = 60
        bar.series[0].graphicalProperties.solidFill = "5B9BD5"
        bar.series[0].graphicalProperties.line.solidFill = "5B9BD5"
        ws.add_chart(bar, f"{get_column_letter(n + 2)}{head + 17}")
    else:
        ws.merge_cells(start_row=head + 1, start_column=1, end_row=head + 1, end_column=n)
        _text(ws.cell(head + 1, 1), "Kỳ này chưa có ngày nào đã qua.")
        ws.cell(head + 1, 1).font = Font(italic=True, color="595959")
    _print_setup(ws, landscape=True)


def build_report(
    schedule: list[ScheduleRow],
    refills: list[ActualRefill],
    days: list[DayRow],
    meta: ReportMeta,
) -> bytes:
    wb = Workbook()
    _sheet_schedule(wb.active, schedule, meta)
    _sheet_refills(wb.create_sheet(), refills, meta)
    _sheet_days(wb.create_sheet(), days, meta)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
